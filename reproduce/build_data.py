"""Rebuild derived display data from preserved official source tables."""
# INTERVIEW: DATA PREPARATION ONLY: this script runs before publication, not in a visitor browser.
# INTERVIEW: It rebuilds derived JSON/CSV from preserved official snapshots. Reading these comments does not require running it.
# INTERVIEW: Do not run build_specs.py over the authored specifications; it is an older presentation generator.
import argparse, json, csv, math, hashlib
from pathlib import Path
from collections import defaultdict

# INTERVIEW: Resolve paths relative to this script: data/source holds inputs, data holds compact chart-ready outputs.
ROOT=Path(__file__).resolve().parents[1]; R=ROOT/'data'/'source'; D=ROOT/'data'
D.mkdir(parents=True,exist_ok=True)
import urllib.request
def read(n): return json.loads((R/(n+'.json')).read_text(encoding='utf-8'))
def save(n,v): (D/(n+'.json')).write_text(json.dumps(v,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
# INTERVIEW: Write each ordinary table twice: JSON for the visualisations, CSV for readers to download.
def table(n,v):
 save(n,v)
 if v:
  with (D/(n+'.csv')).open('w',newline='',encoding='utf-8-sig') as f:
   w=csv.DictWriter(f,fieldnames=list(v[0]));w.writeheader();w.writerows(v)

def make_waterfall(national, market):
 """Decompose the national change without hiding known markets in the residual."""
 totals={r['year']:r['spend'] for r in national if r['traveller']=='International'}
 values={(r['country'],r['year']):r['spend'] for r in market}
 if len(values)!=len(market): raise ValueError('Duplicate market/year records')
 countries=sorted({r['country'] for r in market})
 if any((c,y) not in values for c in countries for y in (2024,2025)):
  raise ValueError('Every market must have both comparison years')
 before,after=totals[2024],totals[2025]
 changes=[(c,values[c,2025]-values[c,2024]) for c in countries]
 drivers=sorted(changes,key=lambda x:(-abs(x[1]),x[0]))[:6]
 drivers=sorted(drivers,key=lambda x:(-x[1],x[0]))
 selected={c for c,_ in drivers}
 # Known omitted markets remain distinct from origins absent from this source table.
 other_published=sum(delta for c,delta in changes if c not in selected)
 unpublished_before=before-sum(values[c,2024] for c in countries)
 unpublished_after=after-sum(values[c,2025] for c in countries)
 unpublished_change=unpublished_after-unpublished_before
 drivers += [('Other published markets',other_published),
             ('Unpublished + rounding',unpublished_change)]
 if sum(delta for _,delta in drivers)!=after-before:
  raise ValueError('Waterfall contributions do not reconcile to national change')
 rows=[dict(label='2024',start=0,end=before/1e9,value=before/1e9,kind='Total',order=0)]
 running=before
 for i,(label,delta) in enumerate(drivers,1):
  rows.append(dict(label=label,start=running/1e9,end=(running+delta)/1e9,
                   value=delta/1e9,kind='Increase' if delta>=0 else 'Decrease',order=i))
  running+=delta
 if running!=after: raise ValueError('Waterfall endpoint does not match national total')
 rows.append(dict(label='2025',start=0,end=after/1e9,value=after/1e9,kind='Total',order=len(rows)))
 return rows

# A bounded update/check supports chart 10 without rewriting any other display data.
parser=argparse.ArgumentParser(description=__doc__)
mode=parser.add_mutually_exclusive_group()
mode.add_argument('--waterfall-only',action='store_true',help='Rebuild only waterfall JSON/CSV from preserved TRA snapshots')
mode.add_argument('--check-waterfall',action='store_true',help='Verify saved waterfall against preserved TRA snapshots without writes')
args=parser.parse_args()
if args.waterfall_only or args.check_waterfall:
 union='Attended or accompanied someone else to business event total'
 national=[dict(year=r['Year'],traveller='International',spend=r['Spend in Australia ($000)']*1000)
           for r in read('tra-summary-dashboard') if r['Category']==union and r['Type']=='International']
 market=[dict(country={'United States of America':'USA','United Kingdom':'UK'}.get(r['Country of origin'],r['Country of origin']),
              year=r['Year'],spend=r['Value']*1000)
         for r in read('tra-international-markets') if r['Attribute']=='Spend in Australia ($000)']
 waterfall=make_waterfall(national,market)
 if args.waterfall_only:
  table('waterfall',waterfall)
 else:
  saved=json.loads((D/'waterfall.json').read_text(encoding='utf-8'))
  if saved!=waterfall: raise ValueError('Saved waterfall JSON differs from source-derived values')
  with (D/'waterfall.csv').open(encoding='utf-8-sig',newline='') as f:
   csv_rows=list(csv.DictReader(f))
  if len(csv_rows)!=len(waterfall): raise ValueError('Waterfall CSV row count differs')
  for actual,expected in zip(csv_rows,waterfall):
   for key,value in expected.items():
    if (actual[key]!=value if isinstance(value,str) else float(actual[key])!=value):
     raise ValueError(f'Waterfall CSV differs at {expected["label"]}: {key}')
  print('Waterfall JSON/CSV match preserved sources; named, other-published and unpublished changes reconcile exactly')
 raise SystemExit(0)

# INTERVIEW: Use preserved Natural Earth inputs when available; download only missing geometry.
for filename,remote in [('world-original.geojson','ne_110m_admin_0_countries.geojson'),('states-original.geojson','ne_50m_admin_1_states_provinces.geojson'),('places-original.geojson','ne_110m_populated_places.geojson')]:
 if not (R/filename).exists(): urllib.request.urlretrieve('https://raw.githubusercontent.com/nvkelso/natural-earth-vector/master/geojson/'+remote,R/filename)

# INTERVIEW: Use the published union of attendees OR accompanying visitors. Adding overlapping subgroup totals would double count.
union='Attended or accompanied someone else to business event total'
summary=[r for r in read('tra-summary-dashboard') if r['Category']==union]
types={'Daytrips':'Domestic day','Overnight':'Domestic overnight','International':'International'}
# INTERVIEW: Prepare national records. Source money is in thousands, so multiply by 1000 to store AUD; spend per trip is spending / trips.
national=[]
for r in summary:
 if r['Type'] not in types:continue
 national.append(dict(year=r['Year'],traveller=types[r['Type']],visitors=r['Visitors'],nights=r['Visitor nights'],spend=r['Spend in Australia ($000)']*1000,total_spend=r['Total trip spend ($000)']*1000,spend_per_trip=r['Spend in Australia ($000)']*1000/r['Visitors']))
# INTERVIEW: assert checks exact reconciliation against the published combined totals; a failed assertion stops this rebuild.
for year in [2024,2025]:
 total=next(r for r in summary if r['Year']==year and r['Type']=='Combined')
 for field,source in [('visitors','Visitors'),('spend','Spend in Australia ($000)'),('total_spend','Total trip spend ($000)')]:
  assert sum(r[field] for r in national if r['year']==year)==total[source]*(1000 if 'spend' in field else 1)
table('national',national)
# INTERVIEW: The domestic presentation uses 2025 only because survey methods changed. Earlier data remain available for international comparisons.
latest=[r for r in national if r['year']==2025]
# INTERVIEW: WAFFLE: calculate each group percentage, floor to whole cells, then distribute remaining cells to the largest fractional remainders.
# INTERVIEW: This largest-remainder allocation ensures each grid contains exactly 100 cells, while share retains the unrounded fraction.
waffle=[]
for metric,key in [('Visitor trips','visitors'),('Spending in Australia','spend')]:
 total=sum(r[key] for r in latest); shares=[r[key]/total*100 for r in latest]
 counts=[math.floor(s) for s in shares]
 for i in sorted(range(3),key=lambda i:shares[i]-counts[i],reverse=True)[:100-sum(counts)]:counts[i]+=1
 index=0
 for r,pct,count in zip(latest,shares,counts):
  for _ in range(count):
# INTERVIEW: x = index modulo 10 gives the column; y counts rows from the other end. These coordinates position each square.
   waffle.append(dict(metric=metric,traveller=r['traveller'],share=pct/100,x=index%10,y=9-index//10));index+=1
 assert index==100
table('waffle',waffle)

# Reuse the bounded narrative transform, so a full rebuild preserves current data.
from story_data import build_outputs, save_outputs
save_outputs(build_outputs())

# INTERVIEW: GEOGRAPHIC JOIN: standardise state names into common codes so TRA, ABS and Natural Earth refer to the same jurisdictions.
abbr={'New South Wales':'NSW','Victoria':'VIC','Queensland':'QLD','South Australia':'SA','Western Australia':'WA','Tasmania':'TAS','Northern Territory':'NT','ACT':'ACT','Australian Capital Territory':'ACT'}
popcols=dict(zip(['NSW','VIC','QLD','SA','WA','TAS','NT','ACT'],range(20,28)))
# INTERVIEW: Read the preserved ABS spreadsheet: row 187 must represent June 2025, aligned with the tourism reference year.
import openpyxl
wb=openpyxl.load_workbook(R/'abs-population-mar-2026.xlsx',read_only=True,data_only=True);ws=wb['Data1']
population={k:ws.cell(187,col).value for k,col in popcols.items()}
assert str(ws.cell(187,1).value).startswith('2025-06')
states=[]
geometry=json.loads((R/'states-original.geojson').read_text(encoding='utf-8'))
features=[]
# INTERVIEW: Round geographic coordinates to reduce file size; this changes display geometry precision, not the tourism estimates.
def round_coords(v):return [round_coords(i) for i in v] if isinstance(v,list) else round(v,3) if isinstance(v,float) else v
for feature in geometry['features']:
 p=feature['properties'];name=p.get('name')
 if p.get('adm0_a3')!='AUS' or name not in abbr:continue
 code=abbr[name]
 candidates=[r for r in read('tra-state-summary') if r['Year']==2025 and r['Type']=='Total' and r['Category']=='Total' and abbr.get(r['State/Territory'])==code]
 assert len(candidates)==1,(name,candidates)
 r=candidates[0];assert r['Sample In/Out']==1
# INTERVIEW: Attach population and spending-per-resident to each state record. Monetary millions are converted to AUD before division.
 item=dict(state=name,code=code,visitors=r['Visitors (000)']*1000,spend=r['Expenditure ($M)']*1e6,nights=r['Nights (000)']*1000,population=population[code],spend_per_resident=r['Expenditure ($M)']*1e6/population[code],longitude=p['longitude'],latitude=p['latitude'])
 # ACT is too small to find without a clearly labelled callout.
 item['label_lon']=153 if code=='ACT' else p['longitude'];item['label_lat']=-36.7 if code=='ACT' else p['latitude']
# INTERVIEW: Embed the joined measures into GeoJSON properties. The choropleth reads properties.spend_per_resident directly; no runtime lookup is needed.
 states.append(item);features.append(dict(type='Feature',properties=item,geometry=dict(type=feature['geometry']['type'],coordinates=round_coords(feature['geometry']['coordinates']))))
assert len(states)==8
table('states',states);save('australia',dict(type='FeatureCollection',features=features))

# INTERVIEW: TREEMAP DATA: root Australia -> states -> destination regions. Parent value=0 avoids adding parent totals to their children again.
regions=[dict(id='Australia',parent=None,name='Australia',value=0)]
for code in popcols:regions.append(dict(id=code,parent='Australia',name=code,value=0))
for r in read('tra-cap-city-regional'):
 if r['Year']!=2025 or r['Region ID'] not in [1,2] or r['Sample In/Out']!=1:continue
 code=abbr.get(r['State/Territory'])
 if not code:continue
 regions.append(dict(id=r['Region'],parent=code,name=r['Region'].replace('Canberra - Regional ACT','ACT (all)').replace(' and the South',' & south').replace('Destination Perth','Perth region'),value=r['Expenditure ($M)']*1e6))
table('regions',regions)

# INTERVIEW: Reshape the source into one row per origin and year, merging its visitor/night/spending measures with a (country, year) key.
markets={}
metric_names={'Visitors (000)':'visitors','Visitor nights (000)':'nights','Spend in Australia ($000)':'spend','Total trip spend ($000)':'total_spend'}
for r in read('tra-international-markets'):
 country={'United States of America':'USA','United Kingdom':'UK'}.get(r['Country of origin'],r['Country of origin'])
 key=(country,r['Year']);markets.setdefault(key,dict(country=key[0],year=key[1]))
 markets[key][metric_names[r['Attribute']]]=r['Value']*1000
market=list(markets.values());assert len(market)==36
table('markets',market)
# INTERVIEW: HEATMAP DATA: one row per origin and measure. Store change as a fraction (0.10 means 10%), plus before/after values.
changes=[]
for country in sorted(set(r['country'] for r in market)):
 before=markets[(country,2024)];after=markets[(country,2025)]
 for key,label in [('visitors','Visitors'),('nights','Nights'),('spend','Spend')]:
  changes.append(dict(country=country,metric=label,change=after[key]/before[key]-1,before=before[key],after=after[key]))
table('changes',changes)
# INTERVIEW: WATERFALL: show six named changes, remaining published markets, then the change
# in national spending not represented by those published markets. The last category includes
# unpublished origins and rounding; its size must not be interpreted as rounding error alone.
# Accumulate whole AUD and convert to billions only for display, preserving exact reconciliation.
waterfall=make_waterfall(national,market)
table('waterfall',waterfall)
# Same six activities in each year, with 2023 source rounding retained.
from activity_data import build_activities, output_text
activities=build_activities()
for suffix, text in zip(('json','csv'),output_text(activities)):
 (D/f'activities.{suffix}').write_text(text,encoding='utf-8')
# INTERVIEW: PER NIGHT: retain the published rounded comparison between business-event and other international visitors.
night=[]
for r in read('tra-per-person-night-spend'):
 night.append(dict(year=str(r['Year']),other=r['Did not attend or accompany anyone to a business event'],business=r[union]))
table('per-night',night)

# INTERVIEW: WORLD BASEMAP: retain country names and rounded geometry; omit Antarctica from this story map.
world=json.loads((R/'world-original.geojson').read_text(encoding='utf-8'))
save('world',dict(type='FeatureCollection',features=[dict(type='Feature',properties={'name':f['properties']['ADMIN']},geometry=dict(type=f['geometry']['type'],coordinates=round_coords(f['geometry']['coordinates']))) for f in world['features'] if f['properties']['ADMIN']!='Antarctica']))
places=json.loads((R/'places-original.geojson').read_text(encoding='utf-8'))
def location(prefix):
 found=[f for f in places['features'] if f['properties'].get('NAME','').startswith(prefix)]
 assert len(found)==1,(prefix,len(found))
 return found[0]['geometry']['coordinates']
capital={'New Zealand':'Wellington','China':'Beijing','USA':'Washington','Singapore':'Singapore','UK':'London','India':'New Delhi','Japan':'Tokyo','Malaysia':'Kuala Lumpur'}
dest=location('Canberra');origins=[];links=[]
# INTERVIEW: FLOW GEOMETRY: generate 61 points along a spherical great-circle connection between representative capitals.
# INTERVIEW: This is a schematic connection, not an observed flight itinerary; the visitor count controls stroke width separately.
def great_circle(a,b):
 def vec(p):
  lon,lat=map(math.radians,p);return [math.cos(lat)*math.cos(lon),math.cos(lat)*math.sin(lon),math.sin(lat)]
 va,vb=vec(a),vec(b);omega=math.acos(sum(x*y for x,y in zip(va,vb)));out=[]
 for i in range(61):
  t=i/60;v=[(math.sin((1-t)*omega)*x+math.sin(t*omega)*y)/math.sin(omega) for x,y in zip(va,vb)]
  out.append([round(math.degrees(math.atan2(v[1],v[0])),3),round(math.degrees(math.atan2(v[2],math.hypot(v[0],v[1]))),3)])
 return out
# INTERVIEW: Take the eight largest published origin markets in 2025 and connect each representative capital to Canberra.
for r in sorted([r for r in market if r['year']==2025],key=lambda r:-r['visitors'])[:8]:
 coord=location(capital[r['country']]);origins.append(dict(country=r['country'],visitors=r['visitors'],longitude=coord[0],latitude=coord[1]))
 links.append(dict(type='Feature',properties={'country':r['country'],'visitors':r['visitors']},geometry={'type':'LineString','coordinates':great_circle(coord,dest)}))
save('routes',dict(type='FeatureCollection',features=links));table('origins',origins)

# INTERVIEW: Record source provenance and processing decisions so the displayed data can be traced back to the official snapshots.
sources=[dict(id='TRA',title='Tourism Research Australia: Business events data',url='https://www.tra.gov.au/en/tourism-statistics/business-events-data',reference='Calendar years 2024 and 2025; public report refreshed 21 May 2026',retrieved='2026-10-05',licence='Commonwealth of Australia, CC BY 4.0',transform='Selected attendance-or-accompaniment union totals; monetary source units converted to AUD; domestic comparisons excluded across 2025 method break. International total spending split into Australian and overseas amounts. State traveller shares use 2025 Category Total components, normalised within each jurisdiction; State Total all-tourism rows are excluded.'),dict(id='ABS',title='ABS National, state and territory population, March 2026',url='https://www.abs.gov.au/statistics/people/population/national-state-and-territory-population/mar-2026',reference='June 2025 population from Table 4, Data1 row 187 columns T:AA',retrieved='2026-10-05',licence='Commonwealth of Australia, CC BY 4.0',transform='Joined eight jurisdictions to TRA spending; spending divided by residents.'),dict(id='NE',title='Natural Earth vector map data',url='https://www.naturalearthdata.com/',reference='1:110m countries and populated places; 1:50m Australian states',retrieved='2026-10-05',licence='Public domain',transform='Selected features; coordinates rounded to three decimals; representative capitals used for schematic flows.')]
save('provenance',sources)
checks=dict(national_reconciled=True,states_joined=len(states),market_rows=len(market),waffle_cells=len(waffle),waterfall_reconciled=True,data_bytes=sum(p.stat().st_size for p in D.glob('*')),excluded=['Spending items: inconsistent aggregate with headline','Days at events: anomalous row and incompatible bins'],caveats=['State visits count a multi-state trip more than once.','Survey estimates and published rounding; no claim of statistical significance.','Amounts are nominal AUD, not inflation adjusted.','2024/2025 domestic data are not directly comparable.'])
(ROOT/'data-checks.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
print(json.dumps(checks,indent=2))

# Rebuild the ranked origin-spending Sankey from the same preserved markets.
from origin_spending import build as build_origin_spending
build_origin_spending()
